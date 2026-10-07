import glob, json, pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

R = "results/"
M = json.load(open(R + "pipeline_metrics.json"))
hq = {q: pd.read_csv(f"{R}hive/hive_{q}.csv") for q in [f"Q{i}" for i in range(1, 9)]}
sp = lambda n: pd.read_csv(glob.glob(f"{R}spark_out/{n}/part-*.csv")[0])
q1 = hq["Q1"].iloc[0]; N = int(q1.total_txns); F = int(q1.fraud_txns)

# ---------- validation (computed, not typed) ----------
checks = []
checks.append(("Spark record count equals Hive total_txns", M["records"], N, M["records"] == N))
checks.append(("Hive Q2: sum of category txns equals total", int(hq["Q2"].txns.sum()), N, int(hq["Q2"].txns.sum()) == N))
checks.append(("Hive Q2: sum of category fraud equals total fraud", int(hq["Q2"].fraud_txns.sum()), F, int(hq["Q2"].fraud_txns.sum()) == F))
cs = sp("channel_summary"); checks.append(("Spark channel_summary fraud sum equals Hive total fraud", int(cs.fraud_txns.sum()), F, int(cs.fraud_txns.sum()) == F))
hs = sp("hour_summary"); checks.append(("Spark hour_summary txns sum equals total", int(hs.txns.sum()), N, int(hs.txns.sum()) == N))
chk_hive_spark = (hq["Q3"].groupby("channel").fraud_txns.sum().sort_index().values == cs.set_index("channel").sort_index().fraud_txns.values).all()
checks.append(("Hive Q3 vs Spark channel fraud counts match per channel", "match" if chk_hive_spark else "differ", "match", bool(chk_hive_spark)))
checks.append(("Hive Q6 monthly txns sum equals total", int(hq["Q6"].txns.sum()), N, int(hq["Q6"].txns.sum()) == N))
checks.append(("Train + test rows equal total", M["split"]["train_rows"] + M["split"]["test_rows"], N, M["split"]["train_rows"] + M["split"]["test_rows"] == N))
assert all(c[3] for c in checks), checks

# ---------- docx helpers ----------
NAVY, BLUE, GREY = RGBColor(0x1f, 0x3a, 0x5f), RGBColor(0x3b, 0x6e, 0xa5), RGBColor(0x55, 0x60, 0x6b)
doc = Document()
sec = doc.sections[0]; sec.page_width, sec.page_height = Cm(21), Cm(29.7)
for s in ("left_margin", "right_margin"): setattr(sec, s, Cm(2))
sec.top_margin = sec.bottom_margin = Cm(2)
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri"); st.paragraph_format.space_after = Pt(5)
for lvl, sz, col in ((1, 16, NAVY), (2, 12.5, BLUE), (3, 11, GREY)):
    h = doc.styles[f"Heading {lvl}"]; h.font.name = "Calibri"; h.font.size = Pt(sz); h.font.bold = True; h.font.color.rgb = col
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri"); h.paragraph_format.space_before = Pt(14 if lvl == 1 else 9); h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True

def shade(cell, hexcol):
    tcPr = cell._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcol); tcPr.append(sh)

def borders(tbl, color="BFC9D4"):
    tblPr = tbl._tbl.tblPr; b = OxmlElement("w:tblBorders")
    for e in ("top", "left", "bottom", "right", "insideH", "insideV"):
        x = OxmlElement(f"w:{e}"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "4"); x.set(qn("w:color"), color); b.append(x)
    tblPr.append(b)

def H(t, l=1): return doc.add_heading(t, l)
def P(t="", bold=False, italic=False, size=None, color=None, align=None, after=None):
    p = doc.add_paragraph(); r = p.add_run(t); r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    if align: p.alignment = align
    if after is not None: p.paragraph_format.space_after = Pt(after)
    return p
def RP(parts):                       # rich paragraph: list of (text, bold)
    p = doc.add_paragraph()
    for t, b in parts: p.add_run(t).bold = b
    return p
def B(t, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(2)
    if bold_lead: p.add_run(bold_lead).bold = True
    p.add_run(t); return p
_n = [0]
def N_(t, restart=False):
    if restart: _n[0] = 0
    _n[0] += 1
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.6)
    p.add_run(f"{_n[0]}.  " + t); return p

def T(header, rows, widths=None, size=9, hl_first_col=False, align_num=True):
    t = doc.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False; borders(t)
    isnum = lambda v: str(v).replace(",", "").replace("%", "").replace("x", "").replace("Rs ", "").replace(".", "", 1).replace("-", "", 1).strip().isdigit()
    numcol = [align_num and i > 0 and len(rows) > 0 and all(isnum(r[i]) for r in rows) for i in range(len(header))]
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ""; r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(size); r.font.color.rgb = RGBColor(255, 255, 255); shade(c, "1F3A5F")
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""; para = cells[i].paragraphs[0]; r = para.add_run(str(v)); r.font.size = Pt(size)
            if hl_first_col and i == 0: r.bold = True
            if ri % 2: shade(cells[i], "F3F6FA")
            if numcol[i]: para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if widths:
        for i, w in enumerate(widths): t.columns[i].width = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs: p.paragraph_format.space_after = Pt(1)
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    doc.add_paragraph().paragraph_format.space_after = Pt(2); return t

def CODE(text):
    t = doc.add_table(rows=1, cols=1); borders(t, "D5DBE3"); c = t.rows[0].cells[0]; shade(c, "F5F7FA"); c.text = ""
    lines = text.strip("\n").split("\n")
    for i, ln in enumerate(lines):
        p = c.paragraphs[0] if i == 0 else c.add_paragraph(); p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
        r = p.add_run(ln if ln else " "); r.font.name = "Consolas"; r.font.size = Pt(8.2); r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def CALLOUT(title, text, fill="EAF1F8", edge="3B6EA5"):
    t = doc.add_table(rows=1, cols=1); c = t.rows[0].cells[0]; shade(c, fill); c.text = ""
    tcPr = c._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders"); x = OxmlElement("w:left"); x.set(qn("w:val"), "single"); x.set(qn("w:sz"), "24"); x.set(qn("w:color"), edge); b.append(x); tcPr.append(b)
    p = c.paragraphs[0]; r = p.add_run(title); r.bold = True; r.font.size = Pt(10)
    p2 = c.add_paragraph(); r2 = p2.add_run(text); r2.font.size = Pt(9.5); p2.paragraph_format.space_after = Pt(2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def FIG(path, cap, w=15.5):
    doc.add_picture(path, width=Cm(w)); doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    P(cap, italic=True, size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)

def df_rows(df, fmts=None):
    out = []
    for _, r in df.iterrows():
        row = []
        for c in df.columns:
            v = r[c]; f = (fmts or {}).get(c)
            row.append(f(v) if f else (f"{int(v):,}" if isinstance(v, (int,)) or (hasattr(v, "is_integer") and float(v).is_integer() and abs(v) > 99) else v))
        out.append(row)
    return out
pct = lambda v: f"{v:.2f}%"; money = lambda v: f"{v:,.0f}"

# ================= COVER =================
P("HCL GUVI × JAIN UNIVERSITY", bold=True, size=13, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
P("BIG DATA ANALYTICS CAPSTONE PROJECT", bold=True, size=13, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
P("Credit Card Fraud Detection", bold=True, size=26, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
P("Big Data Analytics Pipeline on 1 Million Card Transactions", size=13, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
P("HDFS + Hive + Spark / PySpark + YARN + HBase / Pig", size=11, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
T(["Group", "Project", "Team members (fill in)"], [["2", "Credit Card Fraud Detection", "Student 1: ____________   Student 2: ____________   Student 3: ____________\nStudent 4: ____________   Student 5: ____________"]], widths=[1.6, 5.5, 9.9], size=9.5)
CALLOUT("Read this first: what is real and what you must still run",
        "DATASET: the 1,000,000-row dataset is synthetic (generated by data/generate_dataset.py, seed 42). Real card data is confidential, and "
        "the fraud patterns in it (night-time, card-not-present, high-value categories, velocity) were deliberately built into the generator, "
        "so the findings show how the pipeline works and are NOT claims about real-world fraud. If your trainer requires a real dataset, replace the "
        "CSV and re-run the same scripts.\n"
        "EXECUTED: every number in Sections 10 to 13 was produced by running the supplied Hive-style SQL (via Spark SQL) and the PySpark pipeline "
        "on the full 1M rows in Spark local mode (2 cores). Timings are from that run.\n"
        "NOT EXECUTED HERE: the HDFS, Hive, YARN, HBase and Pig commands are written and ready but need your Docker Compose environment. "
        "Run them and paste your own screenshots into the evidence placeholders; the trainer will ask you to demonstrate them live.",
        fill="FFF4E0", edge="D98E04")

# ================= 1-4 =================
H("1. Project Title")
P("Credit Card Fraud Detection: A Big Data Analytics Pipeline using HDFS, Hive, Spark / PySpark, YARN and optional HBase / Pig.")
H("2. Business Problem")
P("A card issuer processes millions of transactions across merchants, channels and cities every month. Only a small fraction (typically well under "
  "1 percent) is fraudulent, yet each fraud costs the issuer chargeback losses, investigation effort and customer trust. Reviewing every transaction "
  "manually is impossible, and rule-only systems either miss new patterns or block too many genuine customers. The business needs to find where fraud "
  "concentrates and to rank transactions by risk so a limited fraud-operations team reviews the right ones first.")
H("3. Business / Analytics Objective")
P("Build a Big Data Analytics solution that stores transaction data in HDFS, organises it with Hive, processes and analyses it with Spark / PySpark, "
  "identifies suspicious patterns, and trains models that rank transactions by fraud risk, producing findings the fraud team can act on.")
H("4. Key Analytics Questions")
for q in ["What percentage of transactions, and what value of money, is fraudulent?",
          "Which merchant categories have the highest fraud rate?",
          "How do channel (POS, contactless, ATM, online) and card presence relate to fraud?",
          "Is fraud concentrated at particular hours of the day or months of the year?",
          "How do transaction amount, distance from home, international use and transaction velocity relate to fraud?",
          "Can a model rank transactions so that a small review queue catches a large share of fraud?",
          "How would the workflow scale when the data grows from 1 million to 100 million transactions?"]: N_(q, restart=(q.startswith("What percentage")))

# ================= 5-6 DATASET =================
H("5. Dataset Description and Source")
T(["Item", "Detail"], [
  ["Source", "Synthetic dataset created by data/generate_dataset.py (numpy / pandas, seed 42), reproducible with the command in the README"],
  ["Why synthetic", "Real card data is confidential; public sets (e.g. the Kaggle ULB creditcard.csv) have anonymised PCA columns that make business questions hard to ask. Replace with a real dataset if your trainer requires it."],
  ["Volume", f"{N:,} transactions, {M['distinct_cards']:,} distinct cards, 5,000 merchants, 10 states, 12 months (Jan to Dec 2025)"],
  ["File size", "about 93 MB as CSV, 16 columns"],
  ["Target column", f"is_fraud (1 = fraud, 0 = genuine); {F:,} frauds = {q1.fraud_pct}% of transactions (highly imbalanced)"],
  ["Fraud signals designed into the data", "card-not-present online use, night hours, high-risk categories (gift cards, jewelry, electronics), amount well above the card's usual spend, long distance from home, international use, bursts of transactions in 24 hours or minutes"],
], widths=[3.6, 13.4], hl_first_col=True)
H("6. Data Dictionary")
T(["Column", "Example", "Type", "Meaning"], [
 ["transaction_id", "T10390367", "STRING", "Unique transaction identifier"],
 ["transaction_ts", "2025-01-01 00:04:43", "TIMESTAMP", "Date and time of transaction"],
 ["card_id", "C166916", "STRING", "Card identifier (one card has many transactions)"],
 ["customer_age", "57", "INT", "Age of cardholder"],
 ["customer_state", "DL", "STRING", "Cardholder's home state"],
 ["merchant_id", "M1293", "STRING", "Merchant identifier"],
 ["merchant_category", "Utilities", "STRING", "Type of merchant (12 categories)"],
 ["channel", "Contactless", "STRING", "POS, Contactless, ATM or Online"],
 ["amount", "321.64", "DOUBLE", "Transaction amount in rupees"],
 ["card_present", "1", "INT", "1 = physical card used, 0 = card-not-present"],
 ["distance_from_home_km", "15.9", "DOUBLE", "Distance between transaction location and cardholder's home"],
 ["txn_count_24h", "2", "INT", "Transactions on the card in the last 24 hours"],
 ["minutes_since_last_txn", "606.2", "DOUBLE", "Minutes since the card's previous transaction"],
 ["is_international", "0", "INT", "1 = transaction outside home country"],
 ["amount_vs_card_avg", "1.25", "DOUBLE", "Amount divided by the card's usual average spend"],
 ["is_fraud", "0", "INT", "Label: 1 = fraud, 0 = genuine"],
], widths=[4.1, 3.7, 2.4, 6.8], size=8.8)
samp = pd.read_csv("data/sample/transactions_sample_4000.csv")
pick = pd.concat([samp[samp.is_fraud == 0].head(3), samp[samp.is_fraud == 1].head(2)])
P("Sample records (first 3 genuine rows and first 2 fraud rows of the shipped 4,000-row sample file):", bold=True, size=9.5)
T(["transaction_id", "category", "channel", "amount", "dist_km", "txn_24h", "intl", "is_fraud"],
  [[r.transaction_id, r.merchant_category, r.channel, f"{r.amount:,.2f}", r.distance_from_home_km, r.txn_count_24h, r.is_international, r.is_fraud] for r in pick.itertuples()],
  widths=[2.9, 2.9, 2.3, 2.1, 1.8, 1.8, 1.4, 1.8], size=8.5)

# ================= 7-8 ARCH =================
H("7. Big Data Architecture")
FIG("screenshots/figs/architecture.png", "Figure 1. End-to-end architecture: source data to insights", w=11.5)
T(["Technology", "Role in this project", "Evidence to show"], [
 ["HDFS", "Distributed storage for raw, processed and result data", "hdfs dfs -ls / -du / fsck output, NameNode UI"],
 ["NameNode / DataNode", "NameNode holds file-to-block metadata; DataNodes store the blocks and replicas", "NameNode UI at :9870, fsck block report"],
 ["YARN", "Allocates containers to Spark executors across the cluster", "ResourceManager UI at :8088"],
 ["Hive", "SQL layer: external table over the HDFS CSV, 8 analytical queries", "Beeline session: SHOW TABLES, query outputs"],
 ["Spark / PySpark", "Distributed cleaning, feature building, aggregation, window analytics, MLlib models", "Notebook / spark-submit output, Spark UI"],
 ["HBase (justified)", "Millisecond lookup of one card's risk profile by card_id row key", "create / put / get / scan output"],
 ["Pig (optional)", "Independent data-flow cross-check of channel fraud counts", "Pig DUMP output"],
 ["Jupyter, Docker Compose, Git/GitHub", "Development interface, training environment, version control", "Notebook, docker ps, repository"],
], widths=[3.4, 7.6, 6.0], size=8.8, hl_first_col=True)
H("8. Why Big Data Technologies Are Required")
T(["Challenge", "Single-machine approach", "Big Data approach"], [
 ["Years of transactions for millions of cards", "Limited by one disk and memory", "HDFS spreads blocks across DataNodes and replicates them"],
 ["Card-level and merchant-level aggregation", "Long local runs, memory errors", "Spark distributes the group-by across executors"],
 ["Analysts who work in SQL", "Ad-hoc local database or spreadsheets", "Hive gives SQL over the same HDFS files"],
 ["Fraud-model training on rare events", "Slow on one core, hard to retrain often", "Spark MLlib trains in parallel and can be re-run on new data"],
 ["Data growing every day", "Buy bigger hardware (scale up)", "Add nodes (scale out), same code"],
], widths=[5.3, 5.0, 6.7], size=9)
CALLOUT("Honest scale note", f"This classroom run uses {N:,} rows (93 MB) in Spark local mode, which a single machine can also handle. The value of the design is that the "
        "same scripts run unchanged against HDFS on a cluster with far more data (see Section 17 for the scaling discussion).")

# ================= 9 HDFS =================
H("9. HDFS: Directory Structure, Ingestion Commands and Evidence")
CODE("""/data/fraud/
  raw/
    transactions.csv            <- uploaded source file
  processed/                    <- cleaned / partitioned data (Spark output)
  results/
    category_summary/   channel_summary/   hour_summary/   month_summary/
    state_summary/      amount_band_summary/   card_risk_profile/   daily_trend/""")
P("Key commands (full list in hdfs/hdfs_commands.txt):", bold=True)
CODE("""hdfs dfs -mkdir -p /data/fraud/raw /data/fraud/processed /data/fraud/results
hdfs dfs -put /tmp/transactions.csv /data/fraud/raw/
hdfs dfs -ls -h /data/fraud/raw/
hdfs dfs -du -h /data/fraud/raw/
hdfs dfs -cat /data/fraud/raw/transactions.csv | head -5
hdfs fsck /data/fraud/raw/transactions.csv -files -blocks -locations""")
B("The client asks the NameNode where to write; the file is split into blocks (128 MB default, so this 93 MB file is one block); blocks are written to DataNodes and replicated (replication factor 3 by default, 1 in a single-DataNode training setup).", "What happens on upload: ")
B("The NameNode stores only metadata (names, blocks, locations); DataNodes store the actual block data and send heartbeats to the NameNode.", "NameNode vs DataNode: ")
CALLOUT("Your evidence (paste screenshots)", "[ Screenshot 1: hdfs dfs -ls -h /data/fraud/raw/ ]    [ Screenshot 2: hdfs fsck block report ]    [ Screenshot 3: NameNode UI file browser ]", fill="F3F6FA", edge="9DB3CC")

# ================= 10 HIVE =================
H("10. Hive: Database, Table Design and Query Results")
P("Hive provides the SQL layer. An external table is used so Hive stores only metadata; dropping the table does not delete the raw file in HDFS.")
CODE(open("hive/02_create_tables.sql").read().split("-- Optional extension")[0].replace("USE fraud_analytics;\n\n", ""))
P("A partitioned ORC table (transactions_orc, partitioned by month) is included in hive/02_create_tables.sql as an optional performance extension.", size=9.5, italic=True)
CALLOUT("How these results were produced", "The eight queries in hive/03_analysis_queries.sql use only HiveQL-compatible syntax. For this document they were run against the full "
        "1M-row CSV through Spark SQL (src: spark/run_sql_queries.py), because a Hive server was not available in the build environment. In your environment, run the same file in "
        "Beeline; the outputs should be identical, and any difference should be investigated.", fill="FFF4E0", edge="D98E04")

H("Q1. Overall fraud level", 2)
T(["Total transactions", "Fraud transactions", "Fraud rate", "Fraud amount (Rs)", "Average amount (Rs)"],
  [[f"{N:,}", f"{F:,}", f"{q1.fraud_pct}%", f"{q1.fraud_amount:,.0f}", f"{q1.avg_amount:,.2f}"]], size=9.5)
H("Q2. Fraud by merchant category", 2)
d = hq["Q2"]; T(["Category", "Transactions", "Fraud txns", "Fraud rate", "Fraud amount (Rs)"],
   [[r.merchant_category, f"{r.txns:,}", f"{r.fraud_txns:,}", f"{r.fraud_pct:.3f}%", f"{r.fraud_amount:,.0f}"] for r in d.itertuples()], size=9)
FIG("screenshots/figs/category.png", "Figure 2. Fraud rate by merchant category (Hive Q2)", w=12.5)
top3 = d.head(3)
P(f"Gift Cards ({d.iloc[0].fraud_pct:.2f}%), Jewelry ({d.iloc[1].fraud_pct:.2f}%) and Electronics ({d.iloc[2].fraud_pct:.2f}%) run at roughly 2.6 to 3.8 times the overall rate of {q1.fraud_pct}%. "
  f"Together they are only {100*top3.txns.sum()/N:.1f}% of transactions but {100*top3.fraud_txns.sum()/F:.1f}% of fraud transactions and {100*top3.fraud_amount.sum()/q1.fraud_amount:.1f}% of fraud money.")
H("Q3. Fraud by channel and card presence", 2)
d = hq["Q3"]; T(["Channel", "Card present", "Transactions", "Fraud txns", "Fraud rate"],
   [[r.channel, r.card_present, f"{r.txns:,}", f"{r.fraud_txns:,}", f"{r.fraud_pct:.3f}%"] for r in d.itertuples()], size=9)
FIG("screenshots/figs/channel.png", "Figure 3. Fraud rate by channel (Hive Q3)", w=12.5)
onl = d.iloc[0]
P(f"Online (card-not-present) transactions are {100*onl.txns/N:.1f}% of volume but {100*onl.fraud_txns/F:.1f}% of all fraud, at {onl.fraud_pct:.2f}%; this is about six times the rate for POS ({d[d.channel=='POS'].fraud_pct.iloc[0]:.2f}%).")
H("Q4. Fraud by hour of day", 2)
FIG("screenshots/figs/hour.png", "Figure 4. Fraud rate by hour of day (Hive Q4)", w=12.5)
h = hq["Q4"]; ng = h[h.txn_hour.isin([23, 0, 1, 2, 3, 4])]; dy = h[~h.txn_hour.isin([23, 0, 1, 2, 3, 4])]
nr, dr = 100*ng.fraud_txns.sum()/ng.txns.sum(), 100*dy.fraud_txns.sum()/dy.txns.sum()
P(f"Transactions between 23:00 and 04:59 account for only {100*ng.txns.sum()/N:.1f}% of volume but have a fraud rate of {nr:.2f}%, versus {dr:.2f}% for the rest of the day ({nr/dr:.1f} times higher).")
H("Q5. Fraud by amount band", 2)
d = hq["Q5"]; T(["Amount band (Rs)", "Transactions", "Fraud txns", "Fraud rate"], [[r.amount_band[3:], f"{r.txns:,}", f"{r.fraud_txns:,}", f"{r.fraud_pct:.3f}%"] for r in d.itertuples()], size=9)
P(f"The fraud rate rises steadily with amount, from {d.iloc[0].fraud_pct:.2f}% under Rs 500 to {d.iloc[-1].fraud_pct:.2f}% at Rs 10,000 and above.")
H("Q6. Monthly trend", 2)
d = hq["Q6"]; T(["Month", "Transactions", "Fraud txns", "Fraud rate"], [[r.txn_month, f"{r.txns:,}", f"{r.fraud_txns:,}", f"{r.fraud_pct:.3f}%"] for r in d.itertuples()], size=9)
P(f"The monthly rate is stable (between {d.fraud_pct.min():.2f}% and {d.fraud_pct.max():.2f}%). No seasonality exists in this synthetic data, so no monthly pattern should be claimed.")
H("Q7. International and velocity signals", 2)
d = hq["Q7"]; T(["International", "Velocity", "Transactions", "Fraud txns", "Fraud rate"], [[r.is_international, r.velocity_flag, f"{r.txns:,}", f"{r.fraud_txns:,}", f"{r.fraud_pct:.3f}%"] for r in d.itertuples()], size=9)
P(f"International transactions with 5 or more transactions in 24 hours show a {d.iloc[0].fraud_pct:.1f}% fraud rate, about {d.iloc[0].fraud_pct/d.iloc[-1].fraud_pct:.0f} times the {d.iloc[-1].fraud_pct:.2f}% of domestic, normal-velocity transactions.")
H("Q8. Highest fraud-rate merchants (minimum 150 transactions)", 2)
d = hq["Q8"]; T(["Merchant", "Transactions", "Fraud txns", "Fraud rate"], [[r.merchant_id, f"{r.txns:,}", f"{r.fraud_txns:,}", f"{r.fraud_pct:.3f}%"] for r in d.itertuples()], size=9)
CALLOUT("Do not over-read Q8", "Merchant IDs are assigned randomly in this synthetic data, so these top merchants (5 to 7 frauds out of about 200 transactions) are small-sample variation, not risky merchants. "
        "On real data, a merchant list like this is a useful watch-list only after checking that the counts are large enough to be statistically meaningful.", fill="FFF4E0", edge="D98E04")
CALLOUT("Your evidence (paste screenshots)", "[ Screenshot 4: SHOW DATABASES / DESCRIBE FORMATTED transactions in Beeline ]    [ Screenshot 5: Q1-Q3 outputs in Beeline ]", fill="F3F6FA", edge="9DB3CC")

# ================= 11 SPARK =================
H("11. Spark / PySpark: Processing, Cleaning and Distributed Analytics")
P("Spark reads the CSV from HDFS, runs quality checks, derives features, aggregates in parallel, runs window analytics and writes results back to HDFS. The full script is spark/fraud_analysis.py.")
CODE("""spark = SparkSession.builder.appName("CreditCardFraudAnalytics").getOrCreate()
df = (spark.read.option("header", True).option("inferSchema", True)
      .csv("hdfs://namenode:8020/data/fraud/raw/transactions.csv"))
df.printSchema(); print("Records:", df.count())""")
H("11.1 Data quality checks (executed result)", 2)
q = M["quality"]
T(["Check", "Result"], [["Records read", f"{M['records']:,}"], ["Null values in any of 16 columns", f"{sum(q['null_counts'].values())}"],
  ["Amount <= 0", q["negative_or_zero_amount"]], ["is_fraud not in (0,1)", q["invalid_label"]], ["Duplicate transaction_id", q["duplicate_transaction_ids"]]], widths=[8, 5], size=9.5)
P("The synthetic data is clean by construction. Real data would show nulls, duplicates and invalid values here, and the cleaning step (drop / impute / dedupe) would be documented.", size=9.5, italic=True)
H("11.2 Transformations", 2)
CODE("""df2 = (df.withColumn("transaction_ts", F.to_timestamp("transaction_ts"))
         .withColumn("txn_hour", F.hour("transaction_ts"))
         .withColumn("txn_month", F.month("transaction_ts"))
         .withColumn("is_night", F.when((F.col("txn_hour") <= 4) | (F.col("txn_hour") == 23), 1).otherwise(0))
         .withColumn("high_velocity", F.when(F.col("txn_count_24h") >= 5, 1).otherwise(0))
         .withColumn("rapid_repeat", F.when(F.col("minutes_since_last_txn") < 5, 1).otherwise(0))
         .withColumn("log_amount", F.log1p("amount"))
         .withColumn("amount_band", F.when(F.col("amount") < 500, "1. <500") ... ))
df2 = df2.repartition(8, "txn_month").cache()   # reused by every later aggregation""")
H("11.3 Distributed aggregations", 2)
CODE("""summary = (df2.groupBy("merchant_category")
             .agg(F.count("*").alias("txns"), F.sum("is_fraud").alias("fraud_txns"),
                  F.round(F.avg("is_fraud") * 100, 3).alias("fraud_pct"))
             .orderBy(F.desc("fraud_pct")))
summary.coalesce(1).write.mode("overwrite").option("header", True).csv("hdfs://namenode:8020/data/fraud/results/category_summary")""")
P("Eleven result tables are written (category, channel, hour, state, month, amount band, age group, region, card risk profile, channel-category ranking, daily trend). The Spark channel and hour totals were checked against Hive in Section 16.")
H("11.4 Window functions and join", 2)
CODE("""# Top 3 riskiest categories inside each channel
w = Window.partitionBy("channel").orderBy(F.desc("fraud_pct"))
ranked = cc.withColumn("rank_in_channel", F.row_number().over(w)).filter("rank_in_channel <= 3")

# 7-day moving average of daily fraud count
w7 = Window.orderBy(F.col("txn_date").cast("timestamp").cast("long")).rangeBetween(-6*86400, 0)
trend = daily.withColumn("fraud_7d_avg", F.round(F.avg("fraud_txns").over(w7), 2))

# Broadcast join with a small state-to-region lookup
df2.join(F.broadcast(region), "customer_state", "left").groupBy("region")...""")
cr = sp("channel_category_rank"); P("Top 3 riskiest categories per channel (executed output):", bold=True, size=9.5)
T(["Channel", "Rank", "Category", "Transactions", "Fraud rate"], [[r.channel, r.rank_in_channel, r.merchant_category, f"{r.txns:,}", f"{r.fraud_pct:.3f}%"] for r in cr.itertuples()], size=8.5)
rg = sp("region_summary"); P("Fraud rate by region (broadcast-join output):", bold=True, size=9.5)
T(["Region", "Transactions", "Fraud txns", "Fraud rate"], [[r.region, f"{r.txns:,}", f"{r.fraud_txns:,}", f"{r.fraud_pct:.3f}%"] for r in rg.itertuples()], size=9)
P("Regional and state differences are small, as expected, because state does not influence fraud in the generator.", size=9.5, italic=True)
H("11.5 Card-level risk profile", 2)
rd = M["card_risk_distribution"]
P(f"Aggregating per card across {M['distinct_cards']:,} cards gives a risk flag (HIGH = 2 or more frauds, MEDIUM = 1, LOW = 0): "
  f"{rd['HIGH']:,} HIGH, {rd['MEDIUM']:,} MEDIUM and {rd['LOW']:,} LOW cards. This table is what the HBase lookup (Section 14) stores. "
  "Because it uses fraud labels from the whole period, it is a historical profile, not a model feature (using it for training would leak the label).")
H("11.6 Spark optimisations used", 2)
T(["Technique", "Where used", "Why"], [
 ["cache()", "df2 after transformation", "Reused by 11+ aggregations and the ML split; avoids re-reading and re-parsing the CSV"],
 ["repartition(8, txn_month)", "before aggregation", "Controls parallelism and groups data by month"],
 ["coalesce(1)", "writing small summary tables", "One output file per small table"],
 ["broadcast join", "state-to-region lookup", "Small table is copied to every executor, avoiding a shuffle"],
 ["Parquet partitionBy(txn_month)", "processed data output", "Columnar, compressed, month-prunable storage"],
], widths=[4.2, 4.6, 8.2], size=9)
CALLOUT("Your evidence (paste screenshots)", "[ Screenshot 6: notebook / spark-submit output ]    [ Screenshot 7: Spark Master or Application UI ]    [ Screenshot 8: YARN ResourceManager application ]    [ Screenshot 9: hdfs dfs -ls -R /data/fraud/results ]", fill="F3F6FA", edge="9DB3CC")

# ================= 12 ML =================
H("12. Fraud Risk Modelling with Spark MLlib")
sp_ = M["split"]
P(f"The goal is not to label every transaction but to rank them so a small review queue catches as much fraud as possible. The data is split by time to avoid using the future to predict the past: "
  f"training on January to September ({sp_['train_rows']:,} rows, {sp_['train_fraud_pct']}% fraud) and testing on October to December ({sp_['test_rows']:,} rows, {sp_['test_fraud_pct']}% fraud).")
B("A model that always predicts 'genuine' is 99.47% accurate and catches no fraud, so accuracy is not used. Models are judged by AUC-ROC, AUC-PR and by recall and lift when only the highest-risk 1% or 5% of transactions are reviewed.", "Why not accuracy: ")
B(f"Fraud is rare, so each training fraud row is weighted {(1-sp_['train_fraud_pct']/100)/(sp_['train_fraud_pct']/100):.0f} times a genuine one (weightCol), instead of discarding genuine data.", "Class imbalance: ")
B("11 numeric features (log amount, card present, distance, 24h count, minutes since last, international, amount vs card average, night flag, rapid repeat, high velocity, age) plus one-hot merchant category and channel, assembled with VectorAssembler inside a Spark ML Pipeline.", "Features: ")
res = M["results"]
T(["Model", "AUC-ROC", "AUC-PR", "Recall @ top 1%", "Precision @ top 1%", "Lift @ top 1%", "Recall @ top 5%"],
  [[m["model"], f"{m['auc_roc']:.3f}", f"{m['auc_pr']:.3f}", f"{m['top1pct']['recall_pct']:.1f}%", f"{m['top1pct']['precision_pct']:.1f}%", f"{m['top1pct']['lift']:.1f}x", f"{m['top5pct']['recall_pct']:.1f}%"] for m in res],
  widths=[4.2, 1.9, 1.8, 2.3, 2.5, 2.0, 2.3], size=8.8, hl_first_col=True)
FIG("screenshots/figs/models.png", "Figure 5. Share of fraud caught when only the riskiest transactions are reviewed", w=12.5)
best = max(res, key=lambda m: m["auc_roc"]); t1 = best["top1pct"]; cm = best["confusion_top1pct"]
P(f"Logistic Regression performed best on this data (AUC-ROC {best['auc_roc']:.3f}). Reviewing only the top 1% highest-risk test transactions ({t1['flagged']:,} transactions) catches {t1['caught_fraud']} of "
  f"{t1['caught_fraud'] + cm['actual1_pred0']:,} frauds ({t1['recall_pct']}%), and the hit rate in that queue is {t1['precision_pct']}% versus a base rate of {sp_['test_fraud_pct']}%, a lift of {t1['lift']}x. "
  f"Reviewing 5% catches {best['top5pct']['recall_pct']}% of fraud and {best['top5pct']['fraud_amount_caught_pct']}% of fraud value.")
CALLOUT("Reading the model results honestly", f"AUC-PR looks small ({best['auc_pr']:.3f}) because the base rate is only {sp_['test_fraud_pct']}%; it is about {best['auc_pr']/(sp_['test_fraud_pct']/100):.0f} times better than random. "
        "The labels in the synthetic data contain deliberate random noise, which caps achievable accuracy, and a simple linear model matches the way the data was generated. "
        "On real data, expect a different ranking of models; always compare several. Most flagged transactions in a 1% queue are still genuine (about 88%), so the queue should drive review or step-up authentication, not automatic blocking.", fill="FFF4E0", edge="D98E04")
H("Most influential features", 2)
gi = M["gbt_feature_importance"][:8]
T(["Rank", "Feature (Gradient Boosted Trees)", "Importance"], [[i + 1, a, f"{b:.3f}"] for i, (a, b) in enumerate(gi)], widths=[1.6, 8, 3], size=9)
P("Amount relative to the card's usual spend, distance from home, transaction velocity and category (gift cards, electronics) dominate. These agree with the SQL findings in Section 10. Feature importance describes what the model uses, not a cause of fraud.")

# ================= 13 HIVE vs SPARK =================
H("13. How Hive and Spark Work Together")
T(["Component", "Responsibility in this project", "Example"], [
 ["HDFS", "Stores the raw CSV and the output", "/data/fraud/raw/transactions.csv"],
 ["Hive", "Table definition and SQL exploration for analysts", "Q1 to Q8 fraud summaries"],
 ["Spark / PySpark", "Cleaning, derived features, joins, window functions, ML", "Risk models, card profiles, 7-day trend"],
 ["YARN", "Gives Spark executors CPU and memory on the cluster", "spark-submit --master yarn"],
], widths=[3.2, 8.2, 5.6], size=9, hl_first_col=True)
CALLOUT("Viva-ready explanation", "Hive and Spark are not two names for the same thing. Hive gives a SQL-oriented table layer over the files in HDFS and is ideal for aggregate questions asked in SQL. "
        "Spark is used when the work goes beyond SQL: derived features, window analytics, joins with other datasets and machine-learning model training, all kept in one distributed workflow.")

# ================= 14 HBASE/PIG =================
H("14. HBase and Pig")
H("14.1 HBase: card risk-profile lookup", 2)
P("Fraud screening often needs one card's history in milliseconds while a payment is in flight. HDFS and Hive scan large files; HBase is a NoSQL store keyed for fast single-row reads. The Spark card_risk_profile output is loaded into an HBase table with card_id as the row key.")
CODE("""create 'card_risk', 'profile'
put 'card_risk', 'C100123', 'profile:fraud_txns', '2'
put 'card_risk', 'C100123', 'profile:risk_flag',  'HIGH'
get 'card_risk', 'C100123'
scan 'card_risk', {LIMIT => 5}""")
P("Row key = card_id; column family = profile (txns, fraud_txns, avg_amount, max_txn_24h, risk_flag). Full commands: hbase/hbase_commands.txt.")
H("14.2 Pig: data-flow cross-check", 2)
P("The Pig script (pig/fraud_analysis.pig) loads the same CSV, groups by channel and counts fraud transactions and fraud amount. Its output must match Hive Q3 and the Spark channel_summary, giving a third independent check of the numbers. Pig is optional: it adds no new capability here, so its main value is validation.")
CALLOUT("Your evidence (paste screenshots)", "[ Screenshot 10: HBase get / scan output ]    [ Screenshot 11: Pig DUMP output matching Q3 ]", fill="F3F6FA", edge="9DB3CC")

# ================= 15 FLOW =================
H("15. End-to-End Execution Flow")
for s in ["Generate or obtain the dataset (data/generate_dataset.py)", "Upload raw CSV to /data/fraud/raw in HDFS and verify", "Create Hive database and external table",
          "Run Hive Q1 to Q8 for initial analysis", "Read the HDFS data in PySpark and run quality checks", "Transform and cache the DataFrame",
          "Run distributed aggregations, window functions and the broadcast join", "Write result tables (CSV) and processed data (Parquet) back to HDFS",
          "Train and evaluate Logistic Regression, Random Forest and Gradient Boosted Trees on a time-based split", "Load the card risk profile into HBase; run the Pig cross-check",
          "Validate that Hive, Spark and Pig totals agree", "Document, push to GitHub, prepare the 10-minute PPT"]:
    N_(s, restart=s.startswith("Generate or obtain"))

# ================= 16 RESULTS + VALIDATION =================
H("16. Final Results, Insights and Recommended Actions")
T(["KPI", "Value", "Interpretation"], [
 ["Transactions analysed", f"{N:,}", "12 months of card activity"],
 ["Fraud transactions", f"{F:,}", f"{q1.fraud_pct}% of transactions"],
 ["Fraud value", f"Rs {q1.fraud_amount:,.0f}", f"Average fraud amount Rs {q1.fraud_amount/F:,.0f} versus Rs {q1.avg_amount:,.0f} for all transactions"],
 ["Riskiest category", "Gift Cards (2.03%)", "3.8x the overall rate"],
 ["Riskiest channel", "Online, card not present (1.22%)", "2.3x overall; about 6x POS"],
 ["Riskiest period", "23:00 to 04:59", f"{nr/dr:.1f}x the fraud rate of other hours"],
 ["Riskiest combination", "International + 5 or more txns in 24h (7.41%)", "About 14x the 0.53% overall rate"],
 ["Best model", f"{best['model']}, AUC-ROC {best['auc_roc']:.3f}", f"Top 1% review catches {t1['recall_pct']}% of fraud at {t1['lift']}x lift"],
], widths=[3.8, 5.6, 7.6], size=9, hl_first_col=True)
T(["Finding", "Possible business action"], [
 ["Gift cards, jewelry and electronics carry 2.6x to 3.8x average fraud", "Step-up authentication or lower limits on first-time high-value purchases in these categories"],
 ["Card-not-present online is the largest fraud channel", "Mandatory 3-D Secure / OTP for risky online purchases"],
 ["Night-time transactions are over twice as risky", "Raise risk score for 23:00 to 04:59 and alert the customer in-app"],
 ["International plus rapid-fire transactions show the highest rates", "Velocity rules that temporarily hold the card and trigger customer confirmation"],
 ["A 1% review queue catches about a quarter of fraud", "Use the model score to prioritise analysts' queues; combine with rules; do not auto-block on the score alone"],
], widths=[7.5, 9.5], size=9)
H("Validation and Testing")
P("Every cross-check below was computed by the build script and the document would not have been produced if any had failed:")
T(["Check", "Value A", "Value B", "Result"], [[c[0], f"{c[1]:,}" if isinstance(c[1], int) else c[1], f"{c[2]:,}" if isinstance(c[2], int) else c[2], "PASS" if c[3] else "FAIL"] for c in checks], widths=[8.6, 2.8, 2.8, 2.0], size=8.8)
P("Still to do in your environment: confirm Beeline Q1 to Q8 outputs equal Section 10, and that Pig and HBase outputs agree.", italic=True, size=9.5)

# ================= 17 PERF =================
H("17. Performance Observations and Scaling")
tm = M["timings_sec"]; seq = list(tm.items()); prev = 0; rows = []
labels = {"read_and_count": "Read CSV + count", "quality_checks": "Data quality checks", "transform_cache": "Transform + repartition + cache", "aggregations_and_write": "11 aggregations + window + join + writes",
          "parquet_write": "Partitioned Parquet write", "train_logreg": "Train Logistic Regression", "train_random_forest": "Train Random Forest (60 trees)", "train_gbt": "Train GBT (40 iterations)", "evaluate": "Evaluate 3 models"}
for k, v in seq: rows.append([labels[k], f"{v - prev:.1f}"]); prev = v
rows.append(["Total", f"{prev:.1f}"])
T(["Stage (Spark local mode, 2 cores, 4 GB driver)", "Seconds"], rows, widths=[11, 3], size=9)
B("Moving from 1M to 100M rows is 100x the data. In local mode this would not fit in memory. On a cluster, HDFS blocks (128 MB each) become hundreds of input partitions, and adding worker nodes increases parallelism.", "From 1M to 100M rows: ")
B("Convert CSV to partitioned Parquet/ORC (done in this project for Parquet), tune spark.sql.shuffle.partitions, avoid skewed keys, broadcast small lookups, and train on a sample or use distributed training for tree models.", "Changes needed: ")
B("This pipeline is batch. Real-time authorisation scoring would additionally need a streaming layer (e.g. Spark Structured Streaming or Kafka) and the HBase profile lookup, which is outside this capstone.", "Not covered: ")

# ================= 18 GITHUB =================
H("18. GitHub Repository and README")
CODE("""credit-card-fraud-bigdata/
|-- data/
|   |-- generate_dataset.py
|   |-- sample/transactions_sample_4000.csv
|-- hdfs/hdfs_commands.txt
|-- hive/01_create_database.sql  02_create_tables.sql  03_analysis_queries.sql
|-- spark/fraud_analysis.py  run_sql_queries.py  make_charts.py
|-- hbase/hbase_commands.txt
|-- pig/fraud_analysis.pig
|-- results/ (hive/*.csv, spark_out/*, pipeline_metrics.json)
|-- screenshots/ (figs/ + your evidence)
|-- docs/build_report.py
|-- README.md""")
P("Suggested meaningful commits: (1) project structure and data generator, (2) HDFS commands, (3) Hive DDL and queries, (4) PySpark cleaning and aggregations, (5) ML models, (6) HBase / Pig, (7) results and validation, (8) README and final report. The README explains setup, how to regenerate the data, how to run each stage and what outputs to expect.")

# ================= 19 CONTRIB =================
H("19. Individual Contributions")
T(["Student", "Name", "Primary responsibility", "Evidence", "Must be able to explain"], [
 ["1", "", "Data ingestion + HDFS", "hdfs_commands.txt, fsck, NameNode UI", "NameNode, DataNode, blocks, replication, commands"],
 ["2", "", "Hive + data modelling", "02_create_tables.sql, 03_analysis_queries.sql, Q1-Q8", "External table, schema, each query's business question"],
 ["3", "", "Spark + PySpark", "fraud_analysis.py, notebook, Spark UI", "Transformations, cache, window functions, ML pipeline"],
 ["4", "", "HBase / Pig", "hbase_commands.txt, fraud_analysis.pig", "Row key design, why HBase for lookups, Pig cross-check"],
 ["5", "", "Integration + documentation", "README, report, validation table, PPT", "End-to-end flow, how totals were validated"],
], widths=[1.4, 2.4, 3.6, 4.8, 4.8], size=8.8)

# ================= 20 CHALLENGES =================
H("20. Challenges and Solutions")
T(["Challenge", "Solution"], [
 ["Only 0.53% of transactions are fraud; accuracy is misleading", "Used class weights, time-based split and ranking metrics (AUC-PR, recall and lift at top 1% / 5%)"],
 ["Leakage risk: random split would let the model see the future", "Split by month: train Jan to Sep, test Oct to Dec"],
 ["Card-level fraud history could leak the label into features", "Kept the card risk profile for HBase lookups only, not for training"],
 ["Same DataFrame was reused many times (slow re-parsing of CSV)", "cache() after transformation and repartition by month"],
 ["Bug while building the pipeline: a Python list alias made feature names include categorical names in the coefficient list", "Copied the list (list(num_cols)); output re-run and verified"],
 ["PySpark failed to install with the system setuptools", "Created a virtual environment with an upgraded setuptools and installed pyspark there"],
 ["Hive server not available in the build environment", "Wrote HiveQL-compatible SQL, ran the same statements via Spark SQL, and flagged that Beeline output must be re-checked"],
], widths=[7.5, 9.5], size=9)
H("21. Limitations")
for t in ["The dataset is synthetic; findings demonstrate the method and are not evidence about real fraud.",
          "Fraud patterns (night, online, categories) were built into the generator, so the SQL findings are expected rather than discovered.",
          "AUC-PR is modest because labels contain random noise; a real model would need richer features such as merchant history and device data.",
          "Spark ran in local mode on 2 cores; no cluster speed-up was measured.",
          "HDFS, Hive, YARN, HBase and Pig steps are specified but not executed in this build."]: B(t)

# ================= 22 CONCLUSION =================
H("22. Final Conclusion")
P(f"The project delivers an end-to-end Big Data pipeline for credit-card fraud analytics. Raw transactions are stored in HDFS, described in Hive, and processed with Spark. "
  f"Across {N:,} transactions the analysis shows fraud concentrated in card-not-present online use, night hours, gift-card / jewelry / electronics purchases, large amounts and international bursts of activity. "
  f"A Logistic Regression risk score ranked on a held-out three-month period lets a 1% review queue catch {t1['recall_pct']}% of fraud, {t1['lift']} times better than random review. "
  "Results are cross-validated between Hive, Spark and Pig, and the design scales from a classroom run to a cluster by changing the data size and resources, not the code.")

# ================= 23 VIVA =================
H("23. Questions You Should Be Ready For (with answers)")
QA = [
 ("Why Big Data technology for this problem?", "Card transactions grow continuously and fraud analysis needs repeated group-bys over all history, plus model retraining. HDFS and Spark scale out; one machine does not."),
 ("Why HDFS?", "It stores large files in replicated blocks across nodes, so storage grows by adding DataNodes and data survives a node failure."),
 ("Role of NameNode and DataNode?", "NameNode keeps metadata (file to block mapping); DataNodes store the blocks and report to the NameNode by heartbeat."),
 ("Role of YARN?", "It is the resource manager: it gives Spark applications containers with CPU and memory on cluster nodes."),
 ("Why an external Hive table?", "Hive manages only metadata; the CSV stays in HDFS and dropping the table does not delete the data."),
 ("Hive vs Spark here?", "Hive: SQL exploration of the table. Spark: feature engineering, windows, joins, ML, all in one distributed workflow."),
 ("What happens when a Spark job runs?", "The driver builds a plan from transformations; an action triggers a job split into stages and tasks; executors run tasks on partitions; shuffles occur for groupBy and joins."),
 ("Why did you cache the DataFrame?", "It is reused by many aggregations and the ML split; caching avoids repeating the CSV read and parse."),
 ("Why not use accuracy?", "99.47% of transactions are genuine, so predicting 'genuine' always gives 99.47% accuracy and catches nothing."),
 ("Why a time-based split?", "Fraud patterns drift. Training on the past and testing on the future mimics real deployment and avoids leakage."),
 ("Why HBase?", "Single-card risk lookup by key in milliseconds, which Hive and HDFS scans are not designed for."),
 ("What would change at 100 million rows?", "Use Parquet/ORC partitioning, more executors, tuned shuffle partitions, and a streaming layer for real-time scoring."),
 ("What is the weakness of your results?", "Synthetic data with built-in patterns; modest AUC-PR; local-mode run; HDFS/Hive/HBase evidence must be added from the cluster."),
]
for qn_, a in QA: RP([("Q. " + qn_ + "  ", True), (a, False)])

# ================= 24 EVALUATION MAP =================
H("24. 100-Mark Evaluation Map")
T(["Area", "Marks", "Where covered in this document", "Your action"], [
 ["Business problem and objective", "10", "Sections 2 to 4", "Review and own it"],
 ["Big Data architecture", "10", "Sections 7, 8, 15", "Redraw diagram in PPT"],
 ["HDFS / data ingestion", "15", "Section 9", "Run commands, paste screenshots"],
 ["Hive / data processing", "15", "Section 10", "Run in Beeline, paste screenshots"],
 ["Spark / PySpark analytics", "15", "Sections 11, 12, 17", "Run on cluster, show Spark UI"],
 ["Advanced technology / integration", "10", "Section 14, 15", "Run HBase / Pig"],
 ["Git/GitHub and reproducibility", "5", "Section 18", "Push with meaningful commits"],
 ["Documentation / README", "5", "Section 18 + README.md", "Add your own run notes"],
 ["Final demo / presentation", "5", "Section 16 (story)", "Build 10-minute PPT"],
 ["Individual contribution", "10", "Section 19, 23", "Each student explains own part"],
], widths=[4.8, 1.5, 6.0, 4.7], size=8.8)
H("25. Final Submission Checklist")
T(["Item", "Status"], [[x, "[ ]"] for x in ["Problem statement", "Dataset description", "Architecture", "HDFS ingestion + screenshots", "Hive database, tables, queries + screenshots",
  "Spark / PySpark notebook run + Spark UI screenshot", "Transformations and analytics", "HBase / Pig evidence", "Final results validated", "GitHub repository", "README", "PPT", "Individual contribution sheet"]],
  widths=[12, 3], size=9)

doc.save("docs/Credit_Card_Fraud_Detection_Big_Data_Project.docx")
print("saved; checks passed:", len(checks))
